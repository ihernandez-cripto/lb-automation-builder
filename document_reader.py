import docx
import openpyxl

def read_docx(file_path: str) -> str:
    doc = docx.Document(file_path)
    full_text = [para.text for para in doc.paragraphs if para.text.strip()]
    return "\n".join(full_text)

def read_excel_config(file_path: str) -> str:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    summary = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        summary.append(f"--- Hoja: {sheet} ---")
        for row in ws.iter_rows(values_only=True):
            row_str = " | ".join([str(cell) for cell in row if cell is not None])
            if row_str:
                summary.append(row_str)
    return "\n".join(summary)